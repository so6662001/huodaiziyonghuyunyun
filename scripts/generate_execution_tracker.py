"""生成增长落地执行进度跟踪 Excel"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

THIN = Side(border_style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
NORMAL_FONT = Font(name="微软雅黑", size=10)
BOLD_FONT = Font(name="微软雅黑", size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")
SUCCESS_FILL = PatternFill("solid", fgColor="D1FAE5")
WARNING_FILL = PatternFill("solid", fgColor="FEF3C7")
DANGER_FILL = PatternFill("solid", fgColor="FEE2E2")
PARAM_FILL = PatternFill("solid", fgColor="FFF2CC")


def style_header(ws, row, n_cols):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = BORDER


def style_data(ws, start_row, end_row, n_cols):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=n_cols):
        for cell in row:
            if not cell.font.bold:
                cell.font = NORMAL_FONT
            cell.alignment = WRAP
            cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = Workbook()

    # Sheet 0: 总览
    ws = wb.active
    ws.title = "0-使用说明"
    ws["A1"] = "货袋子用户增长落地执行进度跟踪"
    ws["A1"].font = Font(name="微软雅黑", size=18, bold=True, color="1F4E78")
    ws.merge_cells("A1:E1")

    rows = [
        [""],
        ["用途", "周度+月度执行进度跟踪", "", "", ""],
        ["维护", "运营负责人 + 各 Owner 实时填写", "", "", ""],
        ["频率", "每周一更新", "", "", ""],
        [""],
        ["📋 Sheet 说明"],
        ["Sheet 0", "本表（使用说明）"],
        ["Sheet 1", "Phase 1 周度任务（W1-12）"],
        ["Sheet 2", "Phase 2 周度任务（W13-24）"],
        ["Sheet 3", "Phase 3 周度任务（W25-36）"],
        ["Sheet 4", "Phase 4 周度任务（W37-52）"],
        ["Sheet 5", "5 大角色周历"],
        ["Sheet 6", "进度看板（每月填）"],
        ["Sheet 7", "风险信号 + 应对"],
        [""],
        ["📋 状态约定"],
        ["未开始", "灰色 = 还没到"],
        ["进行中", "黄色 = 本周正在做"],
        ["已完成", "绿色 = 已 ✓"],
        ["延迟", "红色 = 落后超 1 周"],
        [""],
        ["📋 红绿灯标准"],
        ["🟢 进度达成 ≥ 90%"],
        ["🟡 进度 70-90%（关注）"],
        ["🟠 进度 50-70%（调整）"],
        ["🔴 进度 < 50%（紧急复盘）"],
    ]
    for r in rows:
        ws.append(r)

    for row_idx in [6, 16, 22]:
        for col in range(1, 6):
            ws.cell(row=row_idx, column=col).font = Font(name="微软雅黑", size=12, bold=True, color="1F4E78")

    set_widths(ws, [16, 32, 14, 14, 14])

    # Sheet 1: Phase 1 周度任务（W1-12）
    ws1 = wb.create_sheet("1-Phase1(M0-M2)")
    ws1["A1"] = "Phase 1 准备期（W1-12）"
    ws1["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    ws1.merge_cells("A1:H1")
    ws1.append([])

    ws1.append(["周次", "月份", "主题", "5 个关键任务", "Owner", "验收", "状态", "实际完成"])
    style_header(ws1, 3, 8)

    p1_tasks = [
        ["W1", "M0", "方案对齐", "全员 Kickoff + 招聘启动 + A 类名单 + CEO 短信 + 邀请函 PDF", "CEO+HR+商务", "5/5 全部完成", "未开始", ""],
        ["W2", "M0", "CEO 拜访唐山", "CEO 拜访 10 家 + 客户成功面试", "CEO+HR", "10 家拜访+2 人 Offer", "未开始", ""],
        ["W3", "M0", "CEO 拜访天津", "CEO 拜访 10 家 + 客户成功上岗 + 邀请函发出", "CEO+COO", "10 家+2 人上岗+30 邀请函", "未开始", ""],
        ["W4", "M0", "CEO 拜访聊城+群分级", "CEO 拜访 5 家 + 群分级 + 内容团队就位 + 10 家签约", "CEO+COO+群运营", "5 家+群分级+3 人到岗+10 签约", "未开始", ""],
        ["W5", "M1", "反向邀请上线", "PRD 评审 + 抖音矩阵 + 群内容日历 + 客户成功培训 + 5 家测试", "全员", "5 项全部完成", "未开始", ""],
        ["W6", "M1", "扩种子", "反向邀请扩 15 家 + 群内容执行 + 第一篇行业内容 + A 类 1 对 1", "运营+商务", "15 家激活", "未开始", ""],
        ["W7", "M1", "团队完善", "反向邀请扩 30 家 + 第二批客户成功招聘 + 群运营组长面试 + 地推 JD", "HR+运营", "30 家+2 人 Offer", "未开始", ""],
        ["W8", "M1", "M1 复盘", "数据复盘 + 反向邀请分析 + 卖家委员会候选 + 团队补齐 + M2 计划", "全员", "月报+50 人名单", "未开始", ""],
        ["W9", "M2", "委员会启动", "21 人候选邀请 + 30 家签约 + 地推上岗 + 5 个 KOL + 100 家激活", "CEO+商务", "全部到位", "未开始", ""],
        ["W10", "M2", "扫园开始", "唐山天津各扫园 50 家 + 群运营铺开 + 抖音稳定", "运营+商务", "100 家覆盖", "未开始", ""],
        ["W11", "M2", "委员会确认", "21 委员确认 + 内测合伙人奖励 + 反向邀请监控 + Pre-A 接触", "CEO+CFO", "21 人确认+Pre-A 接触", "未开始", ""],
        ["W12", "M2", "Phase 1 收官", "M2 数据复盘 + 双引擎 V1 评审 + 委员会大会筹备 + KOL 内容 + M3 全员通气", "全员", "Phase 1 末交付", "未开始", ""],
    ]
    for r in p1_tasks:
        ws1.append(r)

    style_data(ws1, 4, ws1.max_row, 8)
    set_widths(ws1, [8, 8, 16, 42, 16, 22, 12, 22])

    # Sheet 2: Phase 2 周度任务（W13-24）
    ws2 = wb.create_sheet("2-Phase2(M3-M5)")
    ws2["A1"] = "Phase 2 试点期（W13-24）"
    ws2["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    ws2.merge_cells("A1:H1")
    ws2.append([])

    ws2.append(["周次", "月份", "主题", "关键任务", "Owner", "验收", "状态", "实际"])
    style_header(ws2, 3, 8)

    p2_tasks = [
        ["W13", "M3", "双引擎上线 ⭐", "产品灰度 5% + 内测合伙人开通 + 委员会会议 + 反向邀请 150 家", "产品+CEO", "5 项完成", "未开始", ""],
        ["W14", "M3", "客户成功冲刺", "5 人 × 20 家 = 100 家 1 对 1 沟通", "客户成功", "接通率 80%+", "未开始", ""],
        ["W15", "M3", "地推+广告", "唐山天津扫园 100 家/人 + 抖音 Dou+ ¥1 万", "商务+运营", "200 家扫园", "未开始", ""],
        ["W16", "M3", "M3 复盘", "数据复盘 + Go/No-Go + 反向邀请优化 + M4 计划", "全员", "月报+决议", "未开始", ""],
        ["W17", "M4", "试单激励", "试单激励上线，前 20 单完成", "运营", "20 单试单", "未开始", ""],
        ["W18", "M4", "群裂变试点", "10 个种子群主试点带群", "群运营", "10 群签约", "未开始", ""],
        ["W19", "M4", "商会合作", "唐山钢贸商会签约", "CEO+商务", "1 个商会签约", "未开始", ""],
        ["W20", "M4", "M4 复盘", "双引擎 25% 灰度 + 数据评估", "全员", "月报", "未开始", ""],
        ["W21", "M5", "KOL 启动", "3 个 KOL 内容产出", "运营", "3 篇内容", "未开始", ""],
        ["W22", "M5", "合伙人评选", "半年优秀奖励发放", "客户成功", "奖励到账", "未开始", ""],
        ["W23", "M5", "Pre-A 接触", "CEO 拜访 3-5 家投资人", "CEO+CFO", "3-5 家拜访", "未开始", ""],
        ["W24", "M5", "Phase 2 收官", "M5 复盘 + Pre-A Term Sheet + 数据交付", "全员", "Term Sheet 签订", "未开始", ""],
    ]
    for r in p2_tasks:
        ws2.append(r)

    style_data(ws2, 4, ws2.max_row, 8)
    set_widths(ws2, [8, 8, 16, 42, 16, 22, 12, 22])

    # Sheet 3: Phase 3 周度任务
    ws3 = wb.create_sheet("3-Phase3(M6-M8)")
    ws3["A1"] = "Phase 3 放量期（W25-36）"
    ws3["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    ws3.merge_cells("A1:H1")
    ws3.append([])

    ws3.append(["周次", "月份", "主题", "关键任务", "Owner", "验收", "状态", "实际"])
    style_header(ws3, 3, 8)

    p3_tasks = [
        ["W25", "M6", "双引擎全网放开 ⭐⭐⭐", "Feature Flag 全开 + 媒体合作 + 地推扩 4 人 + 群主激励 + 反向邀请发放", "全员", "全网放开", "未开始", ""],
        ["W26", "M6", "群裂变", "50 群主接触 + 20 群签约", "群运营", "20 群签约", "未开始", ""],
        ["W27", "M6", "合伙人高峰", "30 家 30 天复盘 + A 类 CEO 二次拜访 + 推荐奖励 + 行业沙龙", "CEO+客户成功", "10 家二次拜访", "未开始", ""],
        ["W28", "M6", "M6 复盘", "数据收集 + 月报 + M7 计划", "全员", "月报", "未开始", ""],
        ["W29", "M7", "群主带群", "目标新增 50 群", "群运营", "50 群", "未开始", ""],
        ["W30", "M7", "KOL 合作", "抖音 KOL 合作（3 个）", "运营", "3 个 KOL", "未开始", ""],
        ["W31", "M7", "上海钢贸展", "赞助+参展", "运营+CEO", "参展完成", "未开始", ""],
        ["W32", "M7", "M7 复盘", "月报 + 第二次委员会", "全员", "委员会会议", "未开始", ""],
        ["W33", "M8", "双 11 筹备", "钢贸版双 11 筹备", "运营", "方案 v1", "未开始", ""],
        ["W34", "M8", "买家裂变", "买家邀请买家灰度", "产品+运营", "灰度上线", "未开始", ""],
        ["W35", "M8", "客户成功扩团", "客户成功扩到 6 人", "HR", "6 人到岗", "未开始", ""],
        ["W36", "M8", "Phase 3 收官", "月报 + 第三次委员会", "全员", "委员会+月报", "未开始", ""],
    ]
    for r in p3_tasks:
        ws3.append(r)

    style_data(ws3, 4, ws3.max_row, 8)
    set_widths(ws3, [8, 8, 16, 42, 16, 22, 12, 22])

    # Sheet 4: Phase 4 周度任务
    ws4 = wb.create_sheet("4-Phase4(M9-M12)")
    ws4["A1"] = "Phase 4 加速期（W37-52）"
    ws4["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    ws4.merge_cells("A1:H1")
    ws4.append([])

    ws4.append(["周次", "月份", "主题", "关键任务", "Owner", "验收", "状态", "实际"])
    style_header(ws4, 3, 8)

    p4_tasks = [
        ["W37", "M9", "跨区扩展", "环渤海+山东扫园计划", "商务", "计划制定", "未开始", ""],
        ["W38", "M9", "买家裂变全开", "买家邀请买家全网上线", "运营+产品", "全网开通", "未开始", ""],
        ["W39", "M9", "A 轮启动", "CEO 接触 A 轮投资人", "CEO", "3 家深度沟通", "未开始", ""],
        ["W40", "M9", "M9 复盘", "月报 + 数据深度分析", "全员", "月报", "未开始", ""],
        ["W41", "M10", "群矩阵", "冲击 4000 群", "群运营", "4000 群", "未开始", ""],
        ["W42", "M10", "双 11 营销", "营销筹备 + 物料就绪", "运营", "物料就绪", "未开始", ""],
        ["W43", "M10", "委员会", "第四次卖家委员会", "CEO", "会议召开", "未开始", ""],
        ["W44", "M10", "M10 复盘", "月报", "全员", "月报", "未开始", ""],
        ["W45", "M11", "双 11 开启", "双 11 钢贸版正式开启", "运营+CEO", "活动启动", "未开始", ""],
        ["W46", "M11", "营销爆发", "媒体集中曝光", "运营", "媒体报道", "未开始", ""],
        ["W47", "M11", "激励翻倍", "反向邀请激励翻倍（限月）", "运营+商务", "激励发放", "未开始", ""],
        ["W48", "M11", "M11 复盘", "月报", "全员", "月报", "未开始", ""],
        ["W49", "M12", "合伙人年会", "内测合伙人年会（含家属）", "CEO", "年会主办", "未开始", ""],
        ["W50", "M12", "白皮书", "年度行业白皮书发布", "内容+CEO", "白皮书发布", "未开始", ""],
        ["W51", "M12", "媒体冲刺", "年度回顾媒体报道", "运营", "媒体集中", "未开始", ""],
        ["W52", "M12", "Y2 规划", "Y2 全员对齐", "全员", "Y2 OKR 通过", "未开始", ""],
    ]
    for r in p4_tasks:
        ws4.append(r)

    style_data(ws4, 4, ws4.max_row, 8)
    set_widths(ws4, [8, 8, 16, 42, 16, 22, 12, 22])

    # Sheet 5: 5 大角色周历
    ws5 = wb.create_sheet("5-角色周历")
    ws5["A1"] = "5 大核心角色每周日程"
    ws5["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    ws5.merge_cells("A1:G1")
    ws5.append([])

    ws5.append(["角色", "周一", "周二", "周三", "周四", "周五", "关键 KPI"])
    style_header(ws5, 3, 7)

    roles = [
        ["CEO", "OKR Weekly + 战略思考", "外部拜访（卖家/投资人/商会）", "KPI Weekly + 产品/技术", "内部 1 对 1（每人 30 分钟）", "财务 review + 公关 + 下周计划",
         "周亲访 ≥5 卖家 / 月 ≥10 家 / 季 ≥1 峰会 / 30% 外部"],
        ["COO", "OKR + 渠道 ROI", "各 VP 1 对 1", "KPI + 100 家健康度 + 群/客户成功巡查", "组织建设（招聘/培训）", "周报 + 月报准备",
         "每周 ROI / 100 家分层动态 / 月 1 次巡查"],
        ["CMO/运营", "OKR + 上周数据", "抖音/视频号内容审核", "群运营巡查 + 内容主编对齐", "地推团队周会", "月度计划 + 预算调整",
         "5 大引擎 ROI / 内容≥30/月 / S+A 群 +5pp/月 / 地推≥30/月"],
        ["客户成功负责人", "OKR + 上周接通数据", "电话 4-5 家", "电话 4-5 家", "电话 4-5 家", "数据复盘 + 团队培训",
         "通话≥200/人/月 / 接通 80% / 开通 50% / 流失<8%"],
        ["客户成功专员", "团队晨会 + 电话", "1对1+电话", "集体培训+电话", "个人复盘+电话", "周报+数据汇总",
         "日 4-5 家电话 / 接通率 / 开通率"],
    ]
    for r in roles:
        ws5.append(r)

    style_data(ws5, 4, ws5.max_row, 7)
    set_widths(ws5, [16, 22, 22, 26, 22, 22, 30])

    # Sheet 6: 进度看板
    ws6 = wb.create_sheet("6-月度进度看板")
    ws6["A1"] = "月度进度看板（每月填写）"
    ws6["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    ws6.merge_cells("A1:G1")
    ws6.append([])

    ws6.append(["月份", "计划新增卖家", "实际", "完成率", "计划新增买家", "实际", "完成率"])
    style_header(ws6, 3, 7)

    targets = [
        ["M1", 8, "", "", 50, "", ""],
        ["M2", 10, "", "", 80, "", ""],
        ["M3", 20, "", "", 300, "", ""],
        ["M4", 35, "", "", 600, "", ""],
        ["M5", 35, "", "", 800, "", ""],
        ["M6", 40, "", "", 1200, "", ""],
        ["M7", 60, "", "", 1500, "", ""],
        ["M8", 70, "", "", 1800, "", ""],
        ["M9", 80, "", "", 2000, "", ""],
        ["M10", 100, "", "", 1700, "", ""],
        ["M11", 120, "", "", 1500, "", ""],
        ["M12", 130, "", "", 1300, "", ""],
        ["Y1 合计", 708, "", "", 12830, "", ""],
    ]
    for r in targets:
        ws6.append(r)

    style_data(ws6, 4, ws6.max_row, 7)
    for cell in ws6[ws6.max_row]:
        cell.font = BOLD_FONT
        cell.fill = SUCCESS_FILL

    # 留空给用户填写实际数据
    for row_idx in range(4, ws6.max_row + 1):
        for col_idx in [3, 4, 6, 7]:
            ws6.cell(row=row_idx, column=col_idx).fill = PARAM_FILL

    set_widths(ws6, [10, 16, 12, 12, 16, 12, 12])

    # Sheet 7: 风险信号
    ws7 = wb.create_sheet("7-风险信号")
    ws7["A1"] = "5 个关键风险信号 + 应对"
    ws7["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    ws7.merge_cells("A1:E1")
    ws7.append([])

    ws7.append(["#", "信号", "触发条件", "应对", "状态"])
    style_header(ws7, 3, 5)

    risks = [
        [1, "M3 末新增卖家不达预期", "< 15 家（计划 20）", "COO+CMO 复盘 / 砍 ROI<8 渠道 / 加大反向邀请 / CEO 加大 A 类回访", ""],
        [2, "M4 末群运营 NPS 低", "< 50", "内容主编上岗考核 / 群运营 1 对 1 培训 / 砍 C 级群 / 群主带群推迟", ""],
        [3, "M5 末资金紧张", "跑道 < 6 月", "砍 PR+部分内容 / CEO 加速 Pre-A / 客户成功精简到 3 人", ""],
        [4, "M6 末反向邀请 ROI 低", "< 8x", "激励梯度重设计 / 反作弊升级 / 邀请话术优化", ""],
        [5, "M8 末渠道 CAC 超标", "> ¥1500", "全面 ROI 评估 / 砍低 ROI / 加大客户成功", ""],
    ]
    for r in risks:
        ws7.append(r)

    style_data(ws7, 4, ws7.max_row, 5)
    set_widths(ws7, [6, 30, 22, 50, 18])

    output = "docs/12-用户增长终版/增长进度跟踪.xlsx"
    wb.save(output)
    print(f"✓ 已生成 {output}")
    print(f"✓ 8 个 Sheet（使用说明 + 4 阶段任务 + 角色周历 + 进度看板 + 风险信号）")


if __name__ == "__main__":
    main()
